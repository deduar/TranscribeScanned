"""Transcribe scanned handwritten PDF notes into a structured .docx.

Pipeline:
    PDF -> page images -> Claude (via Vercel AI Gateway, OpenAI-compatible)
         -> markdown transcript -> .docx with Heading 1/2/3.

Usage:
    python transcribe.py samples/mit_lec01.pdf samples/mit_lec01.docx
"""

from __future__ import annotations

import argparse
import base64
import io
import os
import sys
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from pdf2image import convert_from_path
from PIL import Image
from docx import Document

TRANSCRIPTION_PROMPT = """You are transcribing a single page of scanned handwritten notes.

Rules — follow them exactly:
1. Reproduce every sentence, bullet, and line break EXACTLY as written. Do not paraphrase, summarize, or reorder.
2. Use Markdown structure to mirror visual hierarchy:
   - `# ` for the top-level title/section
   - `## ` for sub-sections
   - `### ` for sub-sub-sections
   - `- ` for bullets
   - Blank line between paragraphs
3. Any word, symbol, or passage you cannot read with high confidence: write `[illegible]` in its place. Do NOT guess.
4. For math/equations, transcribe them inline in plain text (e.g. `V = IR`, `dI/dt = ...`). If a formula is unreadable, write `[illegible equation]`.
5. For diagrams/figures, insert a line: `[figure: short neutral description]`. Do not invent labels you cannot read.
6. Output ONLY the transcribed markdown. No preamble, no commentary, no code fences."""


def encode_image(img: Image.Image) -> str:
    """PNG-encode a PIL image to base64."""
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("ascii")


def transcribe_page(client: OpenAI, model: str, img: Image.Image, page_num: int) -> str:
    """Send one page image to the model, return markdown transcript."""
    b64 = encode_image(img)
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": TRANSCRIPTION_PROMPT},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{b64}"},
                    },
                ],
            }
        ],
        max_tokens=4096,
    )
    text = resp.choices[0].message.content or ""
    print(f"  page {page_num}: {len(text)} chars", file=sys.stderr)
    return text.strip()


def markdown_to_docx(markdown: str, out_path: Path) -> None:
    """Minimal markdown -> docx mapper: handles headings (#, ##, ###) and bullets (-)."""
    doc = Document()
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.lstrip().startswith("- "):
            doc.add_paragraph(line.lstrip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(out_path)


def main() -> int:
    load_dotenv()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path, help="Input scanned PDF")
    parser.add_argument("docx", type=Path, help="Output .docx path")
    parser.add_argument("--pages", type=str, default=None,
                        help="Page range e.g. '1-3' (1-indexed). Default: all.")
    parser.add_argument("--dpi", type=int, default=200)
    args = parser.parse_args()

    api_key = os.environ.get("AI_GATEWAY_API_KEY")
    if not api_key:
        print("ERROR: AI_GATEWAY_API_KEY not set (see env.example)", file=sys.stderr)
        return 2
    model = os.environ.get("MODEL", "anthropic/claude-opus-4-7")

    client = OpenAI(api_key=api_key, base_url="https://ai-gateway.vercel.sh/v1")

    first = last = None
    if args.pages:
        a, _, b = args.pages.partition("-")
        first = int(a)
        last = int(b) if b else first

    print(f"Rendering {args.pdf} at {args.dpi} dpi...", file=sys.stderr)
    images = convert_from_path(
        str(args.pdf), dpi=args.dpi,
        first_page=first, last_page=last,
    )
    print(f"Got {len(images)} page(s). Model={model}", file=sys.stderr)

    chunks: list[str] = []
    start_page = first or 1
    for i, img in enumerate(images):
        page_num = start_page + i
        md = transcribe_page(client, model, img, page_num)
        chunks.append(md)

    combined = "\n\n".join(chunks)
    args.docx.parent.mkdir(parents=True, exist_ok=True)
    markdown_to_docx(combined, args.docx)

    md_path = args.docx.with_suffix(".md")
    md_path.write_text(combined, encoding="utf-8")
    print(f"Wrote {args.docx} and {md_path}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
